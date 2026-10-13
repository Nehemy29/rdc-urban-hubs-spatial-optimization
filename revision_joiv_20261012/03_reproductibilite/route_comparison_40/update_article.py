from pathlib import Path
from copy import deepcopy
import re,json,shutil
import pandas as pd
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.shared import Pt,Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
B=Path(__file__).resolve().parent;D=B.parent/'completion/final';D.mkdir(exist_ok=True)
O=B/'results';summary=pd.read_csv(O/'routing_comparison_summary.csv');long=pd.read_csv(O/'model_vs_osrm_160_comparisons.csv');dc=pd.read_csv(O/'comparaison_documentaire_40_trajets_completee.csv')
assert len(long)==160 and len(summary)==4 and len(dc)==40
labels={'main_road_only':'Main','main_conditional_Bandundu_ferry':'Main + ferry','main_tracks_road_only':'Tracks','main_tracks_conditional_Bandundu_ferry':'Tracks + ferry'}
d=Document(B/'document_inputs/JOIV_revised_manuscript.docx');body=next(p for p in d.paragraphs if p.text.startswith('A World Bank DRC EASE'));head=next(p for p in d.paragraphs if p.text.startswith('N. Historical'));caption=next(p for p in d.paragraphs if p.text.startswith('TABLE XXII'));tablebase=deepcopy(d.tables[-1]._tbl)
def put(anchor,text,template=body):
 el=deepcopy(template._p)
 for x in list(el):
  if x.tag!=qn('w:pPr'):el.remove(x)
 anchor.addprevious(el);p=Paragraph(el,template._parent);r=p.add_run(text)
 if template.runs:r.font.name=template.runs[0].font.name;r.font.size=template.runs[0].font.size;r.bold=template.runs[0].bold;r.italic=template.runs[0].italic
 return p
def table(anchor,title,headers,rows):
 p=put(anchor,title,caption);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.keep_with_next=True
 for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(8);r.bold=False
 el=deepcopy(tablebase);anchor.addprevious(el);t=Table(el,d._body)
 widths=[1.25,.42,.49,.60,.60] if 'XXIII' in title else [.55,.55,.63,.65,.98]
 t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 # Keep 5 original aligned columns; clear unnecessary source rows.
 while len(t.rows)>len(rows)+1:t._tbl.remove(t.rows[-1]._tr)
 while len(t.rows)<len(rows)+1:t.add_row()
 for cells,vals in zip([t.rows[0].cells]+[r.cells for r in t.rows[1:]],[headers]+rows):
  for c,v in zip(cells,vals):c.text=str(v)
 for i,row in enumerate(t.rows):
  for j,c in enumerate(row.cells):
   c.width=Inches(widths[j])
   for p in c.paragraphs:
    p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.left_indent=Inches(0);p.paragraph_format.right_indent=Inches(0);p.paragraph_format.space_after=Pt(2);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(8);r.bold=i==0
 return t
anchor=next(p for p in d.paragraphs if p.text=='III. RESULTS AND DISCUSSION')._p
put(anchor,'O. Forty Route Comparison and Documentary Audit',head)
put(anchor,'An additional audit uses 40 locality pairs and coordinates supplied in a CSV dated 12 October 2026. The input contains OSRM driving-route distances and request URLs, not GPS observations. The unchanged graph-building functions are applied to the original HOT/OSM road archive, reacquired with the same SHA256 hash as the archived source. Four variants are evaluated: main roads and main roads plus tracks, each without and with the three mapped Bandundu ferry ways under the earlier conditional usability assumption. Forward paths use the input coordinates, nearest graph vertices, and geodesic connectors of at most 5 km at both ends. No missing connection is filled to match a reference. The complete sample, unavailable paths and connector offsets are exported.')
put(anchor,'OSRM selects a fastest driving route under its profile [46], while our graph minimizes geodesic length. Both reuse OSM provenance; server data epoch, original raw responses and waypoint snapping offsets are not recorded in the supplied CSV. Thus differences can reflect graph, profile, cost, endpoint and epoch choices. On finite paths, absolute relative difference is 100|d_model−d_OSRM|/d_OSRM and signed difference is d_model−d_OSRM. Means and medians describe the supplied unweighted finite sample, not regional probabilities or calibrated prediction errors. Unavailable routes are reported separately; the 100-km classification check is conditional on finite comparability. No national extrapolation, confidence interval or population-coverage validation is inferred.')
put(anchor,'The supplied documentary table names three non-OSRM references. Two commercial catalogue values are verified on the cited page [47], with undisclosed measurement method, terminals, route date and underlying cartography. The third is a secondary Scribd reproduction whose original provenance could not be verified; it remains flagged and is excluded from confirmed documentary checks. We do not assume that an advertised distance is an independent surveyed road length. All supplied references are preserved in the supplement, without assigning an observed-value status to simulated or cartographic inputs.')
anchor=next(p for p in d.paragraphs if p.text=='IV. CONCLUSION')._p
put(anchor,'O. Forty Route Disagreements and Limited Documentary Checks',head)
rows=[]
for tag in ['main_road_only','main_conditional_Bandundu_ferry','main_tracks_road_only','main_tracks_conditional_Bandundu_ferry']:
 r=summary[summary.scenario.eq(tag)].iloc[0];rows.append([labels[tag],int(r.finite_comparisons),int(r.unavailable),f'{r.mean_absolute_relative_difference_pct:.2f}',f'{r.median_absolute_relative_difference_pct:.2f}'])
table(anchor,'TABLE XXIII\nMODEL AND OSRM DISAGREEMENTS FOR FORTY SUPPLIED ROUTES',['Network','n','No path','Mean %','Median %'],rows)
put(anchor,'Main roads provide 29 finite comparisons and 11 unavailable paths. Mean absolute relative disagreement with OSRM is 24.72%, but the median is 3.30%; mean signed difference is +40.88 km. Adding the conditional Bandundu ferry leaves the finite count at 29 and reduces the mean to 20.47%, with the median unchanged. Including tracks yields 26 finite comparisons and 14 unavailable paths, with mean/median disagreements of 24.74%/5.88%; its conditional ferry variant gives 20.67%/5.88%. Additional track geometry can change nearest-vertex attachment to a disconnected component, so the attachment rule need not produce monotonic reachability. No finite distances are substituted for the unavailable routes.')
put(anchor,'In the main-road comparison, large discrepancies include Binga–Gemena (476.13 km versus 177.62 km OSRM), Kenge–Bandundu (660.77 versus 248.29 km), and Matadi–Boma (228.01 versus 120.79 km). These cases require route-geometry and junction inspection before operational interpretation. Across all four variants, no 100-km classification disagreement occurs among finite comparisons; none of the unavailable paths has an OSRM reference within 100 km. Only three supplied OSRM routes are within 100 km, so the agreement check has limited threshold coverage. This sample-specific finding does not validate regional grid-cell coverage, and it is not evidence that missing links are unimportant. The supplied routes are not a probability sample of household journeys.')
confirmed=dc[dc.ID.isin(['TR-001','TR-018'])]
rows=[]
for r in confirmed.to_dict('records'):
 rows.append([r['ID'],f"{r['Distance documentaire (km)']:.0f}",f"{r['Distance OSRM (km)']:.2f}",f"{r['Distance modèle (km)']:.2f}",f"{r['Écart absolu relatif modèle documentaire pourcent']:.2f}"])
table(anchor,'TABLE XXIV\nTWO COMMERCIAL DOCUMENTARY DISTANCE CHECKS',['Route ID','Doc km','OSRM km','Model km','Abs diff %'],rows)
put(anchor,'TR-001 denotes Matadi–Boma and TR-018 Kinshasa–Mbanza-Ngungu. The commercial page advertises 120 and 155 km, respectively [47]. Main-road model lengths are 228.01 and 142.53 km, yielding absolute relative differences of 90.01% and 8.04%. OSRM lengths are 120.79 and 147.00 km. The Matadi–Boma discrepancy cautions against a general validated-distance claim. The unverified Kinshasa–Kenge value is retained only in the evidence inventory; no aggregate empirical error statistic is estimated from the two commercial entries. These checks strengthen a documented, limited audit, without replacing measured route or household-access observations.')
for p in d.paragraphs:
 if p.text.startswith('Abstract—'):p.text=p.text.replace('These conditional gains complement earlier network and population sensitivity experiments.','A forty-route comparison yields 29 finite main-road paths with 24.72% mean absolute relative disagreement with OSRM; two commercial references provide limited documentary checks.')
 if p.text.startswith('New urban hubs remain the planning objective'):
  p.text=p.text.replace('The previous input and sensitivity audits remain necessary:','The forty-route audit shows substantial routing discrepancies and unavailable paths; commercial checks remain limited and do not establish calibrated accuracy. The previous input and sensitivity audits remain necessary:')
 if p.text.startswith('The replication package provides scripts'):
  p.text=p.text.replace('The consolidated 12 October 2026 package contains the revised documents and the historical-energy diagnostic.','The consolidated 12 October 2026 package also contains the forty-route inputs, 160 computed scenario rows, source-status audit and reproducible routing-comparison code.')
for text in ['[46] Project OSRM contributors, “OSRM HTTP API documentation,” ver. 5.24.0, Route and Table services. [Online]. Available: https://project-osrm.org/docs/v5.24.0/api/. Accessed: 12 Oct. 2026.','[47] DK Transport, “Nos trajets,” commercial route catalogue, Matadi–Boma and Kinshasa–Mbanza-Ngungu entries. [Online]. Available: https://kayou.app/trajets. Verified online: 12 Oct. 2026. Measurement method and route dates not stated.']:
 p=d.add_paragraph(text);p.paragraph_format.space_after=Pt(3)
 for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(8)
# Remove obsolete mid-bibliography section break so references flow in numeric order.
refs=False
for p in d.paragraphs:
 if p.text=='REFERENCES':refs=True
 if refs:
  for node in p._p.xpath('.//w:sectPr'):node.getparent().remove(node)
from docx.enum.section import WD_SECTION_START
d.add_section(WD_SECTION_START.CONTINUOUS)
d.save(D/'JOIV_revised_manuscript.docx')
reply=Document(B/'document_inputs/JOIV_response_to_reviewers.docx')
for p in reply.paragraphs:
 if p.text.startswith('Author:'):p.text='Author: PAMBOLO MAYAMBO Néhémie. Consolidated revision with supplied-route audit: 12 October 2026.'
 if p.text.startswith('Response: Partially addressed with executed external input and route checks'):
  p.text+=' Sections II.O and III.O now add the supplied 40-route audit. We executed four unchanged graph variants, yielding 160 recorded comparisons; main roads give 29 finite paths, 11 unavailable paths and mean/median absolute relative disagreement of 24.72%/3.30% against OSRM. This is a shared-OSM routing comparison, not independent field validation; OSRM minimizes driving time rather than our length objective. Two commercial distances are verified on the cited page, but their original measurement method is unknown. Matadi–Boma differs by 90.01% and Kinshasa–Mbanza-Ngungu by 8.04% from those advertised lengths. The unverified Scribd reproduction is excluded from confirmed documentary checks. We retain unfavourable discrepancies and missing paths rather than claiming a validated regional outcome. Tables XXIII–XXIV, references [46]–[47] and route_comparison_40/results/ provide the evidence.'
 if p.text.startswith('Location and evidence: Sections II.G'):p.text+=' New audit: Sections II.O and III.O; Tables XXIII–XXIV; all supplied route inputs and per-scenario CSVs.'
 if p.text.startswith('External road and published-route checks have been executed.'):
  p.text+=' The new forty-route comparison and two commercial checks remain limited audits: field distances, true service access and calibrated error distributions are still absent. The finite-sample threshold agreement does not remove these limitations.'
 if p.text.startswith('The accompanying revision package contains'):
  p.text+=' The current package includes the supplied cartographic/documentary files, original graph-function source hash, verified road-archive hash, 160 route comparisons, missing-path status and reproduction instructions. No synthetic distance from the earlier SIM-labelled file is used.'
reply.save(D/'JOIV_response_to_reviewers.docx')
assert len(d.tables)==24 and len(d.inline_shapes)==6
print('Updated manuscript: 24 tables, 6 figures, 47 references; response preserves partial-validation status.')
