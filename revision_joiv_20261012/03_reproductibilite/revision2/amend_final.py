from pathlib import Path
from copy import deepcopy
import json,re
import numpy as np,pandas as pd
from docx import Document
from docx.shared import Pt,Mm,RGBColor
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
B=Path(__file__).resolve().parent;F=B/'final';F.mkdir(exist_ok=True);ROOT=B.parent;O=ROOT/'joiv_revision_repo/revision_outputs'
# Reuse the exact verified JOIV table constructor, without rerunning the prior manuscript.
import ast
s=ast.parse((ROOT/'revision/final/build_final.py').read_text());exec(compile(ast.Module(body=[x for x in s.body if isinstance(x,ast.FunctionDef) and x.name in {'p','head','sub','table','setup','setcols'}],type_ignores=[]),'format_helpers','exec'))
def before(anchor,text,style=None):
 z=anchor.insert_paragraph_before(text,style);z.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 if style:z.paragraph_format.first_line_indent=Mm(0)
 return z
def insert_table(anchor,roman,title,headers,rows,widths):
 tmp=Document();table(tmp,roman,title,headers,rows,widths)
 for el in list(tmp._element.body):
  if el.tag!=qn('w:sectPr'):anchor._p.addprevious(deepcopy(el))
# New method and observed-data check.
d=Document(ROOT/'revision/final/JOIV_revised_manuscript.docx')
method=next(x for x in d.paragraphs if x.text=='III. RESULTS AND DISCUSSION')
before(method,'H. Repeated Stresses and Environmental Grid Comparison','Heading 2')
before(method,'An additional experiment holds the main road classes, five-kilometre demand grid, and five-kilometre connectors constant while testing 50-, 100-, and 150-km coverage thresholds. Five independent realizations jointly remove 1% of road features and move hub coordinates 250 m in random directions; seeds are 20261007–20261011, with coordinate azimuth seeds offset by 100000. A separate unperturbed baseline is retained. Both population products and p = 5, 10, and 15 are evaluated, giving 108 additional exact optimizations. These repeated stresses assess realization dependence; five draws do not estimate tails, a calibrated probability distribution, or separate variance contributions of deletion and displacement.')
before(method,'Candidate screening is recomputed at 10 and 25 km on one regional mask, with an EPSG:6933 origin at zero. Centres inside the mask must be at least 40 km from existing hubs and at most 25 km from a mapped road or 15 km from a mapped waterway. Road input is held constant between grids; river, rail, and aerodrome geometry comes from the original retained multimodal inventory. Demographic potential is recomputed from the complete five-kilometre WorldPop grid as Σᵢ Pᵢ/[1+(dᵢⱼ/100)²], using projected distances in kilometres. Normalization bounds are shared across both grids. The six reference preference weights and uniform weights are compared; neither is claimed to be elicited. No new slope, tenure, infrastructure-capacity, or constructability assessment is supplied.')
before(method,'Native WorldCover 2021 v200 pixels are sampled at candidate centres [24]. Strict screening admits classes 20, 30, 40, and 60. An extreme sensitivity condition also admits tree-cover class 10 to measure dependence on the hard exclusion; it does not classify forest degradation or endorse forest conversion. Results refer to centre classifications, not the complete land-cover composition of a 10- or 25-km square. A finer forest-integrity layer is still needed to distinguish primary from degraded secondary cover.')
before(method,'The official UNEP-WCMC protected-area view of WDPCA, October 2026, is queried for COD [33]. It returns 59 polygon records and one point record. Designated, Inscribed, and Established statuses are admitted to the exclusion mask; Proposed records are not treated as established protections. Overlapping records are unioned. Two rules are compared: centre intersection with protected polygons, and any intersection of the full square grid cell with polygons or a point-only reported-area buffer. The latter buffer has radius √(reported area/π) in the equal-area CRS and is a screening assumption, not a legal boundary. Protected-area raw geometries and attributes are not redistributed; users must obtain them from the producer under its terms.')
before(method,'A further source check transcribes the two observed screen-line counts in JICA Table 2.3.18 (printed p. 2-94): SL A, 311,333 PCU/day; SL B, 335,747 PCU/day [32]. The table also reports JICA’s calibrated traffic-model estimates, 287,878 and 316,870. These are genuine secondary traffic observations and another model’s outputs. The present maximal-covering model does not predict PCU/day. Consequently, no predictive error against these counts is assigned to the present model and JICA’s calibration is not presented as its validation.')
# Replace statements made obsolete by actual new computations.
replacements={
'The study therefore uses explicit perturbations rather than claiming a known completeness percentage.':'The study therefore uses explicit perturbations rather than claiming a known completeness percentage.',
'A single deletion and displacement realization per condition cannot characterize stochastic tail behaviour.':'A single deletion and displacement realization per condition cannot characterize stochastic tail behaviour; the supplementary five-realization experiment provides an additional check but still cannot estimate tails.',
'Road additions, speed uncertainty, operational multimodal links, and environmental candidate-grid changes are not included.':'Road additions, speed uncertainty, and operational multimodal links are not included in the 288-run design. Candidate-grid and ecological-screening sensitivity are evaluated separately in Section II.H.',
'The tree-cover definition, 10-versus-25-km candidate-grid comparison, protected areas, tenure, and cadastral constraints are not newly validated.':'The 10-versus-25-km candidate-grid comparison, extreme tree-exclusion sensitivity, and a current protected-area mask are added in Section II.H. Forest integrity, tenure, and cadastral constraints remain unvalidated.',
'The fixed 100-km threshold deserves alternative service-specific tests before use by an institution.':'Alternative 50- and 150-km thresholds are now tested, but no service-specific empirical threshold is calibrated.',
'No protected-area mask or parcel tenure audit is supplied, and the earlier forest exclusions are not newly recomputed under alternative ecological definitions.':'A current protected-area mask and an extreme tree-exclusion sensitivity test are now supplied; parcel tenure and forest integrity remain unverified.',
'The existing Zenodo v1.0.0 archive (10.5281/zenodo.22287394) predates the revision':'The existing Zenodo v1.0.0 archive (10.5281/zenodo.22287394) predates the revision'
}
for par in d.paragraphs:
 txt=par.text
 for a,b in replacements.items():txt=txt.replace(a,b)
 if txt!=par.text:par.text=txt
# Methods are still exploratory, but new results are stated as computed.
a=next(x for x in d.paragraphs if x.text=='IV. CONCLUSION')
before(a,'G. Additional Scenario and Conservation Results','Heading 2')
r=pd.read_csv(O/'repeated_joint_stress_summary.csv');c=pd.read_csv(O/'candidate_grid_protection_counts.csv');env=pd.read_csv(O/'candidate_grid_environment_summary.csv')
before(a,'All 108 additional covering optimizations report zero solver status and zero MIP gap. Table IX gives the five-realization ranges for p = 10; it holds grid resolution and road classes constant. Differences across coverage thresholds express a policy-distance assumption, while variation within a product and threshold reflects the combined deletion–coordinate stress. These ranges remain descriptors of the chosen stresses, not confidence bounds on actual service coverage.')
rows=[]
for prod in ['WorldPop','GHSPOP']:
 for th in [50,100,150]:
  row=r[(r['product']==prod)&(r.threshold_km==th)&(r.p==10)].iloc[0];rows.append([prod,str(th),f'{row["min"]:.2f}–{row["max"]:.2f}',f'{row["mean"]:.2f}'])
insert_table(a,'IX','Ten-hub coverage in five repeated joint stresses',['Product','Limit km','Range %','Mean %'],rows,[26,17,25,18.3])
before(a,'The common regional mask contains 7,677 ten-kilometre centres and 1,239 twenty-five-kilometre centres. Remoteness and mapped-access rules retain 6,519 and 1,047, respectively. Protected polygons contain 1,044 and 167 of those centres. Testing the whole square footprint instead excludes 1,298 and 268 cells, before land-cover screening (Table X). The point-only reported-area buffer intersects no eligible cell in this region. Polygon overlap is dissolved, so cells are counted once. A difference between centre and footprint screening is not a discrepancy in the source; it follows from the explicitly different spatial tests.')
insert_table(a,'X','Candidate grid and protected polygon intersections',['Grid km','Eligible','Centre mask','Cell mask'],[[str(int(x.grid_km)),str(int(x.access_remoteness_eligible)),str(int(x.protected_centre)),str(int(x.protected_cell_intersection))] for x in c.itertuples()],[16,20,25,25.3])
assert env.unclassified_landcover.eq(0).all(),'Do not present incomplete land-cover samples as a completed comparison'
rows=[]
for res in [10,25]:
 for rule,label in [('strict_tree_exclusion','Strict tree exclusion'),('tree_allowed_sensitivity_only','Tree admitted test')]:
  row=env[(env.grid_km==res)&(env.forest_rule==rule)&(env.protection_rule=='cell_intersection')].iloc[0];rows.append([str(res),label,str(int(row.after_landcover)),str(int(row.after_protected_mask))])
insert_table(a,'XI','Land-cover and full-cell protected-area screening',['Grid km','Tree rule','Land-cover pass','After PA mask'],rows,[14,31,21,20.3])
before(a,'Table XI evaluates the land-cover and protected-area rules jointly. Admitting mapped tree cover is deliberately an extreme diagnostic. It measures the influence of the exclusion without identifying which forest could lawfully or ecologically be converted. The detailed output preserves reference and uniform-weight winners at each resolution and mask rule. Changed candidate centres and conditional ranks must not be interpreted as parcel relocation recommendations. Unlike the existing-hub coverage model, this supplementary screen includes preference weights and still requires consultation, slope and flood analysis, and a land-rights assessment.')
before(a,'The strict full-cell screen retains candidates in all ten provinces at 10 km but in only nine at 25 km; Équateur has none at the coarser resolution. Among the nine common provinces, the geodesic shift between the top-scored 10- and 25-km centres has a median of 212.22 km (maximum 313.70 km) under reference weights and 52.02 km (maximum 264.13 km) under uniform weights. Thus, province-specific winners are not stable under this change of sampling grid. The comparison concerns conditional point-centre sampling and ranking, not evidence that urban sites should be moved by those distances.')
before(a,'The JICA screen-line observations are now included as an external evidence inventory. Their units, domain, and model target prevent a valid direct comparison to a percentage of regional population within a distance threshold. JICA’s estimated counts differ from its observations by −7.53% and −5.62%, calculated from its published numbers; these percentages belong to JICA’s own calibration. They are not errors of the model developed here. The principal empirical-validation objection therefore remains unresolved, despite the successful acquisition of actual secondary traffic observations.')
# Add producer reference, replace old primer with a relevant recent primary sensitivity paper.
for par in d.paragraphs:
 if par.text.startswith('[27] '):par.text='[27] A. Puy, S. Lo Piano, A. Saltelli, and S. A. Levin, “sensobol: An R package to compute variance-based sensitivity indices,” Journal of Statistical Software, vol. 102, no. 5, pp. 1–37, 2022, doi: 10.18637/jss.v102.i05.'
z=p(d,'[33] UNEP-WCMC and IUCN, “Protected Planet: The World Database on Protected and Conserved Areas (WDPCA), protected-area view,” Oct. 2026. [Online]. Available: https://www.protectedplanet.net/ and https://data-gis.unep-wcmc.org/server/rest/services/ProtectedSites/The_World_Database_of_Protected_Areas/FeatureServer. Accessed: 7 Oct. 2026.');z.paragraph_format.first_line_indent=Mm(0)
for run in z.runs:run.font.size=Pt(8)
# Preserve original 237-word abstract; insert added experiments by substitution without overstating validation.
for par in d.paragraphs:
 if par.text.startswith('Abstract—'):
  txt=par.text.replace('A factorial stress test evaluates 96 joint configurations for each of three hub counts, yielding 288 optimizations with zero reported mixed-integer programming gaps.','The analysis also includes 288 factorial optimizations and 108 repeated-stress and threshold runs, together with a 10/25-kilometre candidate-grid comparison using current protected-area masks.').replace('An audit of the earlier synthetic graph also reveals substantial centrality rank changes and does not support a predefined p-median elbow criterion.','The synthetic-graph audit reveals centrality rank changes and does not support a predefined p-median elbow criterion.');par.text=txt
  print('abstract words',len(txt.split())-1)
# New conclusion qualifier.
last=next(x for x in d.paragraphs if x.text.startswith('The audits also change'))
last.text=last.text.replace('No empirical regional flow or field validation has been completed,','An additional 108 optimization runs, alternative grid screens, and an October 2026 protected-area mask strengthen sensitivity analysis. No compatible empirical regional flow or field validation has been completed,')
for par in d.paragraphs:
 if par.style.name=='Heading 2':par.alignment=WD_ALIGN_PARAGRAPH.LEFT
# Prevent a one-letter split in the narrow topology header.
d.tables[6].rows[0].cells[2].text='Comps.'
for run in d.tables[6].rows[0].cells[2].paragraphs[0].runs:run.font.name='Times New Roman';run.font.size=Pt(8);run.bold=True
for par in d.paragraphs:
 if re.match(r'^\[\d+\] ',par.text):
  if par.text.startswith('[19] '):par.text=par.text.rstrip().rstrip('.')
  for run in par.runs:run.font.name='Times New Roman';run.font.size=Pt(8)
# A terminal continuous section balances the final reference columns.
sec=d.add_section(WD_SECTION_START.CONTINUOUS);setcols(sec,2)
d.save(F/'JOIV_revised_manuscript.docx')
# Updated response preserves every original numbered comment, adds precise completed evidence.
resp=Document(ROOT/'revision/final/JOIV_response_to_reviewers.docx')
for par in resp.paragraphs:
 txt=par.text.replace('Revision prepared: 6 October 2026.','Revision extended: 7 October 2026.')
 txt=txt.replace('and some proposed greenfield sensitivity analyses have not been performed.','The new extension also compares 10/25-km greenfield screening grids and extreme tree rules with the October 2026 protected-area view, but does not establish forest integrity or constructability.')
 txt=txt.replace('The JICA Kinshasa April 2019 master-plan report was downloaded and its 2017 survey route inventory examined.','The JICA Kinshasa April 2019 master-plan report and appendices were downloaded; its survey-route inventory and Table 2.3.18 observed screen-line counts were examined. Two genuine observations (311,333 and 335,747 PCU/day) are retained. Our model does not predict PCU/day; no matching validation target is available. JICA’s own estimated counts are not attributed to this study.')
 txt=txt.replace('The complete design has 96 configurations per p and 288 exact runs.','The complete design has 96 configurations per p and 288 exact runs.')
 if txt.startswith('Response: Partially addressed by new joint calculations:'):txt+=' The extension adds 108 exact runs: an unperturbed baseline and five independent combined deletion/displacement realizations, both products, three hub counts, and 50/100/150-km coverage thresholds, at 5-km demand resolution. These tests remain uncalibrated scenario stresses, not confidence intervals or Sobol attribution.'
 if txt.startswith('Response: Population product sensitivity'):
  txt='Response: The population alternatives and radius topology tests are completed as previously described. The extension now also compares complete 10- and 25-km candidate screens on a shared boundary mask with a common zero origin and recomputed complete-grid population potential. Native WorldCover samples compare strict tree exclusion with an extreme tree-admitted diagnostic. The official WDPCA protected-area view for October 2026 is applied with centre and full-cell intersection rules. The two grids yield 6,519 and 1,047 access/remoteness-eligible centres, with 1,298 and 268 full-cell protected-polygon intersections before land-cover screening. These calculations address the previously missing grid and hard-forest-rule sensitivity. They do not distinguish primary from degraded forest, certify tenure or flood safety, or constitute a Sobol decomposition. The strict full-cell screen retains 1,583 ten-kilometre candidates and 241 twenty-five-kilometre candidates. Équateur has no retained candidate on the latter grid. Among nine common provinces, the median top-centre shift is 212.22 km under reference weights and 52.02 km under uniform weights, demonstrating substantial grid sensitivity. Road additions and calibrated uncertainty distributions remain untested.'
 txt=txt.replace('Candidate-grid and ecological sensitivity remain unperformed.','Grid and hard-tree-rule comparisons are now computed; forest-integrity and field validation remain unperformed.')
 txt=txt.replace('Greenfield grid-resolution, forest-definition, protected-area, and tenure tests are not supplied.','Greenfield grid-resolution, extreme tree-rule, and current protected-area tests are now supplied. Forest integrity, tenure, and constructability are not validated.')
 txt=txt.replace(', The new extension', '. The new extension')
 if txt.startswith('Location and evidence: Sections II.B, II.E–F, III.E–F; Table II;'):txt+=' Extension: Sections II.H and III.G, Tables IX–XI.'
 if txt!=par.text:par.text=txt
head(resp,'Additional evidence supplied on 7 October 2026')
p(resp,'New manuscript Section II.H specifies repeated stresses, candidate-grid alignment, the complete-grid population potential, WorldCover sampling, and the current protected-area query. Section III.G and Tables IX–XI report the corresponding results. The new CSV/JSON files include repeated_joint_stress, candidate_grid_environment, candidate_grid_protection_counts, and jica_observed_screenline_benchmark. Raw WDPCA data are withheld from redistribution under producer terms. Their producer query and release are documented for lawful reacquisition.')
p(resp,'The strongest outstanding objection remains empirical validation of the study’s regional accessibility or coverage target. Obtaining an observed traffic table does not validate a model that predicts a different quantity. We therefore do not claim all reviewer requests are satisfied, and respectfully seek evaluation as an exploratory geospatial sensitivity study. If compatible regional validation remains mandatory, that evidence must still be collected or obtained before the stronger claim can be supported.')
resp.save(F/'JOIV_response_to_reviewers.docx')
sup=Document(ROOT/'revision/final/JOIV_supplementary_methods.docx')
head(sup,'S7 Extension executed on 7 October 2026')
for name in ['repeated_joint_stress_metadata.json','candidate_grid_environment_metadata.json']:
 head(sup,name);p(sup,json.dumps(json.loads((O/name).read_text()),indent=2,ensure_ascii=False))
head(sup,'S8 Secondary observed traffic values')
p(sup,'JICA Table 2.3.18, printed p. 2-94, reports observed screen-line counts of 311,333 PCU/day for SL A and 335,747 for SL B. Its calibrated model estimates are 287,878 and 316,870. Calculated relative differences are −7.533734% and −5.622388%. These are JICA model comparisons only. They are not predictions or validation errors of our hub-coverage model. The main report and appendices were inspected; raw GPS trajectories and a compatible regional OD target were not obtained.')
head(sup,'S9 Protected-area inputs and redistribution')
p(sup,'UNEP-WCMC and IUCN (2026), Protected Planet: The World Database on Protected and Conserved Areas (WDPCA), October 2026, Cambridge, UK: UNEP-WCMC and IUCN, https://www.protectedplanet.net/. The public protected-area view is filtered by iso3 containing COD, outSR=4326. It returns 59 polygon records and one point record without transfer-limit truncation. Designated, Inscribed and Established records enter the mask. Polygon union handles overlapping designations. Full-cell and centre tests are distinguished; reported-area point buffers are hypothetical circular screening footprints. Raw geometries and source attributes must not be included in the distributed revision archive. Obtain data from the producer under the current terms. Access date: 7 October 2026.')
p(sup,'SHA-256 source polygon response: d7a8b8ef6b45e01117ab4693a3b5dccdf00f0982f1076ef3c815ef70dfa7a1ac. Point response: 3435b3c82dd5f341fa1be38a9044bc20ec8e5e6a221f010725409795b3142f3e. File hashes preserve the response snapshot; they do not imply redistribution permission.')
for par in sup.paragraphs:
 if 'Revision prepared 6 October 2026' in par.text:par.text=par.text.replace('Revision prepared 6 October 2026','Revision extended 7 October 2026')
 if par.text.startswith('{') or par.text.strip().startswith('{'):
  par.alignment=WD_ALIGN_PARAGRAPH.LEFT;par.paragraph_format.first_line_indent=Mm(0)
  for run in par.runs:run.font.name='Consolas';run.font.size=Pt(8)
 if '288 joint maximal-covering result rows' in par.text:par.text=par.text.replace('288 joint maximal-covering result rows','288 factorial and 108 additional repeated-stress maximal-covering result rows')
sup.save(F/'JOIV_supplementary_methods.docx')
# Usable expert questionnaire, clearly prospective.
q=Document();setup(q);q.sections[0].left_margin=q.sections[0].right_margin=Mm(22);q.styles['Normal'].font.size=Pt(11)
head(q,'Questionnaire de pondération pour le diagnostic territorial')
q.paragraphs[0].style='Title';q.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
q.paragraphs[0].runs[0].font.size=Pt(16)
p(q,'Étude de PAMBOLO MAYAMBO Néhémie — Université de Kinshasa. Ce formulaire prépare une consultation volontaire sur les critères d’un diagnostic territorial. Il ne constitue pas une consultation déjà réalisée. Les réponses servent à comparer des préférences professionnelles ; elles ne certifient pas la faisabilité d’un site.')
head(q,'1 Consentement et profil')
p(q,'Je souhaite participer volontairement et j’accepte l’utilisation agrégée de mes réponses pour cette étude : Oui □  Non □. Aucun nom n’est nécessaire dans le jeu d’analyse. Code attribué par le chercheur : __________. Date : __________.')
p(q,'Domaine principal : aménagement □  transport □  environnement □  géomatique □  économie territoriale □  administration locale □  autre : __________. Années d’expérience : _____. Connaissance du système Grand Bandundu–Grand Équateur–Kinshasa–Kongo Central : faible □  moyenne □  forte □.')
head(q,'2 Répartition de 100 points')
p(q,'Attribuez au total 100 points aux six critères ci-dessous. Davantage de points signifie davantage d’importance pour une présélection territoriale. Cette présélection précède les études foncières, hydrologiques et environnementales. Ne copiez pas une pondération proposée par le chercheur : indiquez votre préférence.')
for criterion,desc in [('Population accessible','Potentiel démographique autour du candidat, calculé sur une grille de population.'),('Proximité routière','Proximité de routes cartographiées, sans garantie de praticabilité.'),('Proximité des voies d’eau','Proximité de cours d’eau cartographiés, sans garantie de navigabilité.'),('Proximité ferroviaire','Proximité de lignes cartographiées, sans garantie de fonctionnement.'),('Proximité d’aérodrome','Proximité d’un aérodrome cartographié, sans garantie de desserte.'),('Éloignement des pôles existants','Possibilité de desservir des espaces éloignés des pôles retenus.')]:
 p(q,criterion+' : ____ points. '+desc)
p(q,'TOTAL : ____ points (doit être égal à 100).')
head(q,'3 Justification et critères absents')
p(q,'Pourquoi avez-vous attribué les deux poids les plus élevés ?\n________________________________________________________________\n________________________________________________________________')
p(q,'Quels critères importants manquent au diagnostic ?\n________________________________________________________________\n________________________________________________________________')
p(q,'Quelles données locales pourraient vérifier ou corriger ces critères ? Indiquez l’institution ou le document, sans transmettre de données personnelles.\n________________________________________________________________\n________________________________________________________________')
head(q,'4 Protocole du chercheur après collecte')
p(q,'Documenter le nombre de participants, leurs profils et le recrutement réellement effectué. Rejeter ou faire corriger les totaux différents de 100. Conserver les réponses brutes anonymisées, puis calculer la moyenne, la médiane et la dispersion de chaque poids. Comparer les classements avec les pondérations individuelles, les poids uniformes et les poids de référence. Signaler la taille et les limites de l’échantillon ; ne pas revendiquer un consensus national. Respecter les exigences institutionnelles applicables à la consultation. Aucun effectif ni avis d’expert ne doit être inventé.')
for par in q.paragraphs:
 par.alignment=WD_ALIGN_PARAGRAPH.LEFT
 if par.text.startswith('3 Justification'):par.paragraph_format.page_break_before=True
 if par.text.startswith(('Pourquoi','Quels critères','Quelles données')):par.text+='\n________________________________________________________________\n________________________________________________________________'
q.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
q.styles['Title'].font.name='Times New Roman';q.styles['Title'].font.color.rgb=RGBColor(0,0,0)
for run in q.paragraphs[0].runs:run.font.name='Times New Roman';run.font.color.rgb=RGBColor(0,0,0);run.font.size=Pt(16)
for node in [q.styles['Title']._element,q.paragraphs[0]._p]:
 for border in node.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
q.save(F/'JOIV_questionnaire_experts_FR.docx')
print('Documents amended; candidate landcover complete; no external validation fabricated')
