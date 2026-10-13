from pathlib import Path
import json,hashlib,zipfile,shutil,importlib.metadata,ast
from docx import Document
import pandas as pd,numpy as np,rasterio
R=Path(__file__).resolve().parent.parent;B=R/'revision2';F=B/'final';repo=R/'joiv_revision_repo';S=repo/'revision_sources';O=repo/'revision_outputs'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
# Preserve the executable winner comparison as part of the candidate script.
p=repo/'scripts/revision_candidate_grid_environment.py';s=p.read_text()
if '# Geodesic comparison of top centres' not in s:
 s+='''
# Geodesic comparison of top centres under strict full-cell screening.
from pyproj import Geod
w=pd.DataFrame(selected);shifts=[]
for weight in ['reference','uniform']:
 a=w[(w.forest_rule=='strict_tree_exclusion')&(w.protection_rule=='cell_intersection')&(w.weight==weight)]
 b=a[a.grid_km==10].merge(a[a.grid_km==25],on='province',suffixes=('_10','_25'))
 for row in b.itertuples():
  distance=Geod(ellps='WGS84').inv(row.longitude_10,row.latitude_10,row.longitude_25,row.latitude_25)[2]/1000
  shifts.append(dict(weight=weight,province=row.province,winner_10km=row.candidate_id_10,winner_25km=row.candidate_id_25,geodesic_shift_km=distance))
sh=pd.DataFrame(shifts);sh.to_csv(O/'candidate_grid_winner_shifts.csv',index=False)
(O/'candidate_grid_winner_shifts_metadata.json').write_text(json.dumps({'strict_fullcell_common_provinces':int(sh[sh.weight=='reference'].province.nunique()),'comparison':sh.groupby('weight').geodesic_shift_km.agg(['count','median','max']).to_dict('index')},indent=2))
'''
 p.write_text(s)
for p in (repo/'scripts').glob('revision_*.py'):ast.parse(p.read_text())
# Verify the new exact solutions and ecological derived results.
t=pd.read_csv(O/'repeated_joint_stress.csv');assert len(t)==108 and t.solver_status.eq(0).all() and t.mip_gap.eq(0).all()
assert t.groupby(['seed','product','p']).apply(lambda z:z.sort_values('threshold_km').coverage_pct.is_monotonic_increasing,include_groups=False).all()
a=pd.read_csv(O/'candidate_grid_comparison.csv');assert len(a)==7566 and a.worldcover_code.notna().all() and not a.candidate_id.duplicated().any()
# Check a native-pixel sample against local source tiles independently of the saved cache.
for tile,g in a.sample(100,random_state=6222).groupby('tile'):
 with rasterio.open(S/'worldcover'/(tile+'.tif')) as ds:
  sample=np.array([x[0] for x in ds.sample(zip(g.longitude,g.latitude))]);assert np.array_equal(sample,g.worldcover_code.to_numpy())
for m in json.loads((S/'worldcover_tile_manifest.json').read_text()):assert sha(S/'worldcover'/(m['tile']+'.tif'))==m['sha256']
# Source descriptors only: never distribute producer protected geometries.
man=json.loads((S/'source_manifest.json').read_text());p=S/'jica_appendices.pdf'
if not any(x['file']==p.name for x in man):man.append(dict(file=p.name,url='https://openjicareport.jica.go.jp/pdf/12340329_01.pdf',version='JICA PDTK April2019 appendices',bytes=p.stat().st_size,sha256=sha(p),access_date='2026-10-07'))
(S/'source_manifest.json').write_text(json.dumps(man,indent=2))
req=(R/'revision/final/requirements_revision.txt').read_text()
# Preserve the original recorded NetworkX audit version; the extension does not import it.
for line in req.splitlines():
 name,version=line.split('==')
 if name!='networkx':assert importlib.metadata.version(name)==version
(F/'requirements_revision.txt').write_text(req)
# Copy only latest, visually checked PDF exports.
for stem,folder in [('JOIV_revised_manuscript','render3/manuscript'),('JOIV_response_to_reviewers','render/JOIV_response_to_reviewers'),('JOIV_supplementary_methods','render2/supplement'),('JOIV_questionnaire_experts_FR','render2/questionnaire')]:shutil.copy2(B/folder/(stem+'.pdf'),F/(stem+'.pdf'))
d=Document(F/'JOIV_revised_manuscript.docx');abstract=next(p.text for p in d.paragraphs if p.text.startswith('Abstract'))
assert len(abstract.split())-1==230 and len(d.tables)==11 and len(d.inline_shapes)==4
assert len([p for p in d.paragraphs if p.text.startswith('[') and '] ' in p.text])==33
report=dict(date='2026-10-07',main_optimization_runs=396,extension_runs=108,extension_all_solver_status_zero=True,extension_all_gap_zero=True,native_worldcover_candidates=7566,unknown_worldcover=0,native_samples_independently_checked=100,protected_records_polygons=59,protected_established_polygons=58,protected_point_records=1,protected_raw_redistributed=False,abstract_words=230,manuscript_tables=11,manuscript_figures=4,manuscript_references=33,remaining=['compatible empirical validation','actual expert elicitation','forest integrity and tenure','calibrated uncertainty distributions and operational routing'],independent_clean_machine_run=False)
(F/'artifact_verification.json').write_text(json.dumps(report,indent=2))
readme=(R/'revision/final/README_REVISION.md').read_text().replace('Prepared 6 October 2026','Extended 7 October 2026')
readme+='''
## Extension of 7 October 2026 — current submission files
The CURRENT manuscript, response, supplement and prospective expert questionnaire are in revision2/final. Files in revision/final are retained as prior inputs to the amendment script, not the current submission versions.

Added results: 108 exact runs (two population products, three hub counts, 50/100/150 km thresholds, five independent combined deletion/displacement draws plus a baseline), for 396 covering optimizations including the original factorial design. These are scenario stresses, not calibrated probability or Sobol sensitivity indices.

The 10/25 km candidate screens use one boundary mask and grid origin, the same mapped access inventory, recomputed full-grid population potential, native WorldCover2021v200 centre samples, and centre versus full-cell protected-area masks. The strict full-cell screens retain 1,583 candidates at 10 km and 241 at 25 km. Équateur has no retained candidate at 25 km. Province-specific winners change substantially; these results do not certify new-city locations. Tree admission is solely an extreme sensitivity test, not a forest-conversion recommendation.

JICA Table2.3.18 observations (311,333 and 335,747 PCU/day) are retained as real secondary evidence. This study does not predict PCU/day: JICA's own calibrated estimates are not validation of our population coverage model. Empirical validation remains unresolved. The expert questionnaire is a prospective collection instrument; no consultation is claimed to have occurred.

Run the extension after the preceding analyses:

    python joiv_revision_repo/scripts/revision_repeated_stress.py
    python joiv_revision_repo/scripts/revision_download_protected.py
    python joiv_revision_repo/scripts/revision_candidate_grid_environment.py
    python revision2/amend_final.py

Protected Planet WDPCA October2026 data are private producer inputs. Raw polygons and source attributes are NOT in this archive. The query script checks the recorded response hashes and stops on changed data; obtain the documented archived snapshot or explicitly disclose and rerun on a newer release. Native WorldCover tiles are acquired on demand by the candidate script and verified against worldcover_tile_manifest.json; the preserved derived sampling cache avoids needless repeated source retrieval. Do not redistribute WDPCA input files in a public release. Cite UNEP-WCMC and IUCN, Protected Planet WDPCA, October2026, and observe producer terms.

The current extension was executed and checked in this environment. A complete clean-machine rerun of all earlier phases has not been performed. The amendment script uses the retained original DOCX inputs; Word rendering is an additional verification step described in the accompanying document workflow. A new GitHub/Zenodo release must archive this revision before it can receive its own release DOI.
'''
(F/'README_REVISION.md').write_text(readme)
# Preserve earlier inputs and replace code/tables; include only a whitelist of metadata.
updates={}
for p in (repo/'scripts').glob('revision_*.py'):updates[str(p.relative_to(R))]=p
for p in O.rglob('*'):
 if p.is_file() and p.suffix in {'.csv','.json','.gz','.png'}:updates[str(p.relative_to(R))]=p
for name in ['source_manifest.json','worldcover_tile_manifest.json','protected_source_metadata.json']:updates['joiv_revision_repo/revision_sources/'+name]=S/name
updates['joiv_revision_repo/data/02_traitees/reseau_multimodal_phase5.gpkg']=repo/'data/02_traitees/reseau_multimodal_phase5.gpkg'
for p in F.iterdir():
 if p.is_file() and p.suffix in {'.docx','.pdf','.json','.md','.txt'}:updates[str(p.relative_to(R))]=p
updates['revision2/amend_final.py']=B/'amend_final.py';updates['revision2/finish_package.py']=B/'finish_package.py'
checks={};archive=F/'JOIV_revision_complete.zip'
with zipfile.ZipFile(R/'recovered/JOIV_revision_complete.zip') as old,zipfile.ZipFile(archive,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=6) as new:
 for entry in old.infolist():
  if entry.filename in updates or entry.filename in {'README.md','CHECKSUMS_SHA256.json'}:continue
  data=old.read(entry.filename);new.writestr(entry.filename,data);checks[entry.filename]=hashlib.sha256(data).hexdigest()
 for name,p in updates.items():new.write(p,name);checks[name]=sha(p)
 new.writestr('README.md',readme);checks['README.md']=hashlib.sha256(readme.encode()).hexdigest()
 new.writestr('CHECKSUMS_SHA256.json',json.dumps(checks,indent=2))
with zipfile.ZipFile(archive) as z:
 assert z.testzip() is None
 assert not any('wdpca_cod_layer' in n or 'wdpa2017' in n or '/worldcover/' in n for n in z.namelist())
 for name,digest in checks.items():assert hashlib.sha256(z.read(name)).hexdigest()==digest
print('Verified',len(checks),'archive entries',archive.stat().st_size,'bytes; 100 native samples checked; 396 runs documented')
