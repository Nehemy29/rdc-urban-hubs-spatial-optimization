from pathlib import Path
from docx import Document
B=Path(__file__).resolve().parent;D=B/'deliverables';D.mkdir(exist_ok=True)
REPO='https://github.com/Nehemy29/rdc-urban-hubs-spatial-optimization'
DOI='https://doi.org/10.5281/zenodo.23389655'
changes={
'JOIV_revised_manuscript.docx':('The replication package provides scripts', 'The replication package is archived as version v2.0.0 on Zenodo ('+DOI+'), published on 15 October 2026 from GitHub commit 8a4a186. It provides scripts, derived results, parameters, source metadata, checksums and reproduction instructions, including the forty-route inputs and 160 computed scenario rows. The repository is '+REPO+'. The original v1.0.0 archive (https://doi.org/10.5281/zenodo.22287395) predates these revisions; https://doi.org/10.5281/zenodo.22287394 identifies all versions. The present submission files add this archive citation after publication; the archived v2.0.0 calculation files and tag remain unchanged. Large raw third-party datasets are reacquired from their producers under documented terms. The small OSM ferry extract and OSM-derived diagnostic paths are supplied with ODbL attribution; protected-area and satellite reference geometries are not redistributed.'),
'JOIV_response_to_reviewers.docx':('The accompanying revision package contains','The revised code and derived results are publicly archived as v2.0.0 on Zenodo ('+DOI+'), published on 15 October 2026 from GitHub commit 8a4a186: '+REPO+'/releases/tag/v2.0.0. The archive contains the revised Python scripts, tables, metadata, figures, supplied cartographic/documentary files, graph-function and road-archive source hashes, 160 route comparisons, missing-path status and reproduction instructions. No synthetic distance from the earlier SIM-labelled file is used. The original v1.0.0 archive has version-specific DOI 10.5281/zenodo.22287395; DOI 10.5281/zenodo.22287394 identifies all versions. The present submission documents add the confirmed archive citation after release; the archived calculations and v2.0.0 tag remain unchanged. Public archiving supports reproducibility and does not close the empirical-validation limitations discussed above.'),
'Mode_emploi_renvoi_JOIV_FR.docx':('Publier les nouveaux scripts','La version GitHub v2.0.0 et son archive Zenodo sont publiées depuis le 15 octobre 2026. DOI à citer pour cette révision : '+DOI+'. Le dépôt est '+REPO+'. Le DOI 10.5281/zenodo.22287395 désigne v1.0.0; 10.5281/zenodo.22287394 désigne toutes les versions. Le manuscrit et la réponse fournis ici incluent le DOI confirmé; ils sont à utiliser pour le renvoi. Le code et les calculs archivés sous v2.0.0, commit 8a4a186, restent figés. La mise à jour de citation du 15 octobre est distincte et ne nécessite pas de refaire les calculs ou de déplacer l’étiquette v2.0.0. La note Zenodo indiquant que le DOI sera ajouté provient du texte antérieur archivé; les métadonnées à jour sont dans le dépôt courant. Les données brutes lourdes et les sources sous restrictions ne sont pas toutes redistribuées.')
}
for name,(prefix,new) in changes.items():
 d=Document(B/'inputs'/name);matches=[p for p in d.paragraphs if p.text.startswith(prefix)];assert len(matches)==1,(name,len(matches));p=matches[0]
 if p.runs:
  p.runs[0].text=new
  for run in p.runs[1:]:run.text=''
 else:p.add_run(new)
 if name=='Mode_emploi_renvoi_JOIV_FR.docx':
  for p in d.paragraphs:
   if p.text=='Actualiser GitHub et Zenodo':p.text='Archive GitHub et Zenodo confirmée'
   if p.text.startswith('PAMBOLO MAYAMBO Néhémie'):p.text='PAMBOLO MAYAMBO Néhémie — citation actualisée le 15 octobre 2026'
 d.save(D/name)
 original=Document(B/'inputs'/name)
 assert len(d.tables)==len(original.tables)
 assert len(d.inline_shapes)==len(original.inline_shapes)
 assert len(d.paragraphs)==len(original.paragraphs)
 print(name,'DOI updated; tables, images and paragraph counts preserved')
